#!/usr/bin/env python3
"""Build the formatted BE Project Work - II Word report.

Usage (from repository root):

    .venv/bin/python report/build_report.py

The generated DOCX follows the attached Stage-II template: A4 margins,
Times New Roman 12 pt justified body text at 1.5 line spacing, 14 pt
centred chapter titles, 12 pt section headings, 10 pt centred footer,
Roman front-matter page numbers and Arabic numbering from Chapter 1.
"""
from __future__ import annotations

import math
import re
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

import report_content as C

ROOT = Path(__file__).resolve().parents[1]
FIGDIR = ROOT / "report" / "figures"
OUT = ROOT / "report" / "HydroRainGuard_BE_Project_Stage_II_Report.docx"

PAGES = C.PAGES
MAIN_PAGE_COUNT = len(PAGES)
REF_START_PAGE = MAIN_PAGE_COUNT + 1
FRONT_LABELS = {
    "certificate": "i",
    "acknowledgement": "ii",
    "abstract": "iii",
    "toc1": "iv",
    "toc2": "v",
    "abbrev": "vi",
    "nomenclature": "vii",
    "lof1": "viii",
    "lof2": "ix",
    "lot1": "x",
    "lot2": "xi",
}


# ---------------------------------------------------------------------------
# Low-level OOXML helpers
# ---------------------------------------------------------------------------

def set_run_font(run, name: str = "Times New Roman", size: float | None = None,
                 bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:cs"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_font(style, name: str = "Times New Roman", size: float = 12,
                   bold: bool = False, italic: bool = False) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:cs"), name)


def add_style(doc: Document, name: str, size: float, *, bold: bool = False,
              italic: bool = False, align=WD_ALIGN_PARAGRAPH.LEFT,
              line_spacing: float = 1.5, first_indent: float = 0.0,
              space_before: float = 0, space_after: float = 0,
              font_name: str = "Times New Roman",
              base: str = "Normal") -> None:
    styles = doc.styles
    try:
        style = styles[name]
    except KeyError:
        style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = styles[base]
    set_style_font(style, font_name, size, bold, italic)
    pf = style.paragraph_format
    pf.alignment = align
    pf.line_spacing = line_spacing
    pf.first_line_indent = Inches(first_indent)
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.widow_control = True


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    set_style_font(normal, "Times New Roman", 12)
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing = 1.5
    pf.first_line_indent = Inches(0.5)
    pf.space_before = Pt(0)
    pf.space_after = Pt(12)
    pf.widow_control = True

    add_style(doc, "ChapterTitle", 14, bold=True,
              align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.5,
              space_before=0, space_after=10)
    doc.styles["ChapterTitle"].paragraph_format.keep_with_next = True

    add_style(doc, "SectionHeading", 12, bold=True,
              align=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.5,
              space_before=8, space_after=6)
    doc.styles["SectionHeading"].paragraph_format.keep_with_next = True

    add_style(doc, "SubHeading", 12, bold=True,
              align=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.5,
              space_before=6, space_after=4)
    doc.styles["SubHeading"].paragraph_format.keep_with_next = True

    add_style(doc, "FrontHeading", 14, bold=True,
              align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.5,
              space_before=0, space_after=12)
    doc.styles["FrontHeading"].paragraph_format.keep_with_next = True

    add_style(doc, "CoverLine", 12, bold=True,
              align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.0,
              space_before=0, space_after=6)
    add_style(doc, "CoverTitle", 16, bold=True,
              align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.0,
              space_before=0, space_after=8)
    add_style(doc, "CertificateBody", 12,
              align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.5,
              first_indent=0.5, space_after=10)
    add_style(doc, "Signature", 12,
              align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.0,
              space_before=4, space_after=4)
    add_style(doc, "Caption", 12,
              align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.0,
              space_before=2, space_after=0)
    add_style(doc, "Legend", 10,
              align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.0,
              space_before=0, space_after=8)
    add_style(doc, "TableCaption", 12, bold=True,
              align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.0,
              space_before=4, space_after=4)
    doc.styles["TableCaption"].paragraph_format.keep_with_next = True
    add_style(doc, "TableCellText", 9,
              align=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.0,
              space_after=0)
    add_style(doc, "TableHeadText", 9, bold=True,
              align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.0,
              space_after=0)
    add_style(doc, "TOCEntry", 10,
              align=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.0,
              space_after=0)
    add_style(doc, "ListEntry", 9,
              align=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.0,
              space_after=1)
    add_style(doc, "CodeBlock", 9, font_name="Courier New",
              align=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.0,
              space_after=0)
    add_style(doc, "ReferenceEntry", 10.5,
              align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.15,
              space_after=7)
    doc.styles["ReferenceEntry"].paragraph_format.left_indent = Inches(0.3)
    doc.styles["ReferenceEntry"].paragraph_format.first_line_indent = Inches(-0.3)
    add_style(doc, "FooterText", 10,
              align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.0,
              space_after=0)
    add_style(doc, "PageNumberText", 10,
              align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.0,
              space_before=0, space_after=0)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top: int = 60, start: int = 80,
                     bottom: int = 60, end: int = 80) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_field(paragraph, instruction: str, placeholder: str = "1") -> None:
    run = paragraph.add_run()
    set_run_font(run, "Times New Roman", 10)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(text)
    run._r.append(fld_end)


def set_page_numbering(section, fmt: str, start: int) -> None:
    sect_pr = section._sectPr
    pg_num_type = sect_pr.find(qn("w:pgNumType"))
    if pg_num_type is None:
        pg_num_type = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num_type)
    pg_num_type.set(qn("w:fmt"), fmt)
    pg_num_type.set(qn("w:start"), str(start))


def set_update_fields(doc: Document) -> None:
    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def configure_section(section, *, fmt: str | None = None, start: int | None = None) -> None:
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(23.4)
    section.left_margin = Mm(37)
    section.bottom_margin = Mm(32)
    section.right_margin = Mm(23.4)
    section.header_distance = Mm(12)
    section.footer_distance = Mm(12)
    if fmt and start is not None:
        set_page_numbering(section, fmt, start)


def configure_footer(section, with_page_number: bool) -> None:
    # Each section needs its own footer part: cover has no page number,
    # front matter uses Roman numerals and the main report uses Arabic.
    section.footer.is_linked_to_previous = False
    footer = section.footer
    for p in list(footer.paragraphs):
        p._element.getparent().remove(p._element)
    p1 = footer.add_paragraph(style="FooterText")
    p1.add_run("RSCOE, Department of Information Technology 2026-27")
    if with_page_number:
        p2 = footer.add_paragraph(style="PageNumberText")
        add_field(p2, " PAGE \\* MERGEFORMAT ", "1")


# ---------------------------------------------------------------------------
# Text / reference handling
# ---------------------------------------------------------------------------

def iter_content_strings(pages: list[dict]):
    for page in pages:
        if page.get("chapter_title"):
            yield page["chapter_title"]
        yield page["heading"]
        yield from page["paras"]
        for f in page.get("figures", []):
            yield f["caption"]
            yield f["legend"]
        if page.get("table"):
            yield page["table"]["caption"]
            for row in page["table"]["rows"]:
                yield from row
        yield from page.get("code", [])


def build_reference_map(pages: list[dict]) -> tuple[dict[str, int], list[str]]:
    order: list[str] = []
    pattern = re.compile(r"\[\[([A-Za-z0-9_]+)\]\]")
    for text in iter_content_strings(pages):
        for key in pattern.findall(str(text)):
            if key not in C.REFS:
                raise KeyError(f"Unknown reference key: {key}")
            if key not in order:
                order.append(key)
    return {key: i + 1 for i, key in enumerate(order)}, order


def replace_refs(text: str, ref_map: dict[str, int]) -> str:
    def repl(m: re.Match[str]) -> str:
        key = m.group(1)
        if key not in ref_map:
            raise KeyError(f"Unknown reference key: {key}")
        return f"[{ref_map[key]}]"
    return re.sub(r"\[\[([A-Za-z0-9_]+)\]\]", repl, str(text))


# ---------------------------------------------------------------------------
# Paragraph / table / figure helpers
# ---------------------------------------------------------------------------

def body_paragraph(doc: Document, text: str, *, style: str = "Normal",
                   page_break: bool = False):
    p = doc.add_paragraph(style=style)
    if page_break:
        p.paragraph_format.page_break_before = True
    p.add_run(text)
    return p


def front_heading(doc: Document, text: str, *, page_break: bool = False):
    p = doc.add_paragraph(style="FrontHeading")
    if page_break:
        p.paragraph_format.page_break_before = True
    p.add_run(text)
    return p


def chapter_title(doc: Document, text: str, *, page_break: bool = False):
    p = doc.add_paragraph(style="ChapterTitle")
    if page_break:
        p.paragraph_format.page_break_before = True
    p.add_run(text.upper())
    return p


def section_heading(doc: Document, text: str, *, page_break: bool = False,
                    style: str = "SectionHeading"):
    p = doc.add_paragraph(style=style)
    if page_break:
        p.paragraph_format.page_break_before = True
    p.add_run(text)
    return p


def spacer(doc: Document, count: int = 1, size: float = 12):
    for _ in range(count):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(" ")
        set_run_font(r, "Times New Roman", size)


def cover_line(doc: Document, text: str, size: float = 12, bold: bool = True,
               style: str = "CoverLine", page_break: bool = False):
    p = doc.add_paragraph(style=style)
    if page_break:
        p.paragraph_format.page_break_before = True
    r = p.add_run(text)
    set_run_font(r, "Times New Roman", size, bold=bold)
    return p


def add_table(doc: Document, rows: list[list[str]], caption: str, number: str,
              ref_map: dict[str, int], *, font_size: float = 9,
              page_break: bool = False) -> None:
    cap = doc.add_paragraph(style="TableCaption")
    if page_break:
        cap.paragraph_format.page_break_before = True
    cap.add_run(f"Table {number}: {replace_refs(caption, ref_map)}")
    ncol = max(len(r) for r in rows)
    table = doc.add_table(rows=1, cols=ncol)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    total_w = 5.85
    col_w = total_w / ncol
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    set_row_cant_split(hdr)
    for i in range(ncol):
        cell = hdr.cells[i]
        cell.width = Inches(col_w)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "E8E8E8")
        set_cell_margins(cell)
        cp = cell.paragraphs[0]
        cp.style = doc.styles["TableHeadText"]
        run = cp.add_run(replace_refs(rows[0][i] if i < len(rows[0]) else "", ref_map))
        set_run_font(run, "Times New Roman", font_size, bold=True)
    for row_vals in rows[1:]:
        row = table.add_row()
        set_row_cant_split(row)
        for i in range(ncol):
            cell = row.cells[i]
            cell.width = Inches(col_w)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell)
            cp = cell.paragraphs[0]
            cp.style = doc.styles["TableCellText"]
            run = cp.add_run(replace_refs(row_vals[i] if i < len(row_vals) else "", ref_map))
            set_run_font(run, "Times New Roman", font_size)
    # A small spacer keeps following captions/legends clear.
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(0)
    sp.paragraph_format.line_spacing = 1.0
    r = sp.add_run(" ")
    set_run_font(r, "Times New Roman", 4)


def add_code_block(doc: Document, lines: list[str], ref_map: dict[str, int],
                   *, page_break: bool = False) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F2F2F2")
    set_cell_margins(cell, 100, 120, 100, 120)
    p = cell.paragraphs[0]
    p.style = doc.styles["CodeBlock"]
    if page_break:
        p.paragraph_format.page_break_before = True
    text = "\n".join(replace_refs(line, ref_map) for line in lines)
    run = p.add_run(text)
    set_run_font(run, "Courier New", 9)
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(0)
    sp.paragraph_format.line_spacing = 1.0
    r = sp.add_run(" ")
    set_run_font(r, "Times New Roman", 4)


def add_figure(doc: Document, spec: dict, number: str, ref_map: dict[str, int],
               *, multi: bool = False, page_break: bool = False) -> None:
    path = FIGDIR / spec["path"]
    if not path.exists():
        raise FileNotFoundError(path)
    with Image.open(path) as im:
        px_w, px_h = im.size
    aspect = px_w / px_h
    max_w = 5.7
    if multi:
        max_h = 2.08
    elif aspect < 1.0:
        max_h = 4.35
    else:
        max_h = 3.00
    w_in = min(max_w, max_h * aspect)
    h_in = w_in / aspect
    if h_in > max_h:
        h_in = max_h
        w_in = h_in * aspect
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.first_line_indent = Inches(0)
    if page_break:
        p.paragraph_format.page_break_before = True
    run = p.add_run()
    run.add_picture(str(path), width=Inches(w_in), height=Inches(h_in))
    cap = doc.add_paragraph(style="Caption")
    cap.add_run(f"Figure {number}: {replace_refs(spec['caption'], ref_map)}")
    leg = doc.add_paragraph(style="Legend")
    leg.add_run(replace_refs(spec["legend"], ref_map))


def add_index_table(doc: Document, headers: list[str], rows: list[list[str]],
                    *, page_break: bool = False, font_size: float = 9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [0.8, 4.35, 0.7] if len(headers) == 3 else [1.4, 1.55, 1.4, 1.55]
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    set_row_cant_split(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.width = Inches(widths[i])
        set_cell_shading(cell, "E8E8E8")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.style = doc.styles["TableHeadText"]
        r = p.add_run(h)
        set_run_font(r, "Times New Roman", font_size, bold=True)
        if page_break and i == 0:
            p.paragraph_format.page_break_before = True
    for row_vals in rows:
        row = table.add_row()
        set_row_cant_split(row)
        for i, val in enumerate(row_vals):
            cell = row.cells[i]
            cell.width = Inches(widths[i])
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.style = doc.styles["ListEntry"]
            r = p.add_run(str(val))
            set_run_font(r, "Times New Roman", font_size, bold=(len(headers) == 3 and str(row_vals[0]).startswith("APPENDIX")))
    return table


# ---------------------------------------------------------------------------
# Front matter
# ---------------------------------------------------------------------------

def add_cover(doc: Document) -> None:
    spacer(doc, 4)
    cover_line(doc, "A PROJECT REPORT ON", 12, True)
    spacer(doc, 2)
    cover_line(doc, "HYDRO RAIN GUARD:", 16, True, "CoverTitle")
    cover_line(doc, "INTELLIGENT RAINWATER HARVESTING AND", 16, True, "CoverTitle")
    cover_line(doc, "ARTIFICIAL GROUNDWATER RECHARGE", 16, True, "CoverTitle")
    cover_line(doc, "ASSESSMENT SYSTEM", 16, True, "CoverTitle")
    spacer(doc, 2)
    cover_line(doc, "SUBMITTED TO THE SAVITRIBAI PHULE PUNE UNIVERSITY, PUNE", 12, True)
    cover_line(doc, "IN THE PARTIAL FULFILLMENT OF THE REQUIREMENTS", 12, True)
    cover_line(doc, "FOR THE AWARD OF THE DEGREE", 12, True)
    spacer(doc, 1)
    cover_line(doc, "OF", 12, True)
    spacer(doc, 1)
    cover_line(doc, "BACHELOR OF TECHNOLOGY", 16, True, "CoverTitle")
    cover_line(doc, "(INFORMATION TECHNOLOGY)", 16, True, "CoverTitle")
    spacer(doc, 1)
    cover_line(doc, "SUBMITTED BY", 12, True)
    spacer(doc, 1)
    cover_line(doc, "MR. SHUBHAM DESAI       ROLL NO: __________  EXAM NO: __________", 11, True)
    cover_line(doc, "MR. SIDDHESH MAHADIK    ROLL NO: __________  EXAM NO: __________", 11, True)
    cover_line(doc, "MR. KAUSHAL AGALE       ROLL NO: __________  EXAM NO: __________", 11, True)
    cover_line(doc, "MR. OMKAR KSHIRSAGAR    ROLL NO: __________  EXAM NO: __________", 11, True)
    spacer(doc, 1)
    cover_line(doc, "DEPARTMENT OF INFORMATION TECHNOLOGY", 16, True, "CoverTitle")
    spacer(doc, 1)
    cover_line(doc, "JSPM'S RAJARSHI SHAHU COLLEGE OF ENGINEERING", 12, True)
    cover_line(doc, "TATHAWADE, PUNE 411033", 12, True)
    cover_line(doc, "(AN AUTONOMOUS INSTITUTE AFFILIATED TO SAVITRIBAI PHULE PUNE UNIVERSITY, PUNE)", 10, True)
    cover_line(doc, "2026-2027", 12, True)


def add_certificate(doc: Document) -> None:
    front_heading(doc, "CERTIFICATE")
    spacer(doc, 2)
    p = body_paragraph(doc, "This is to certify that the project report entitled", style="CertificateBody")
    p.paragraph_format.first_line_indent = Inches(0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cover_line(doc, "\u201C" + C.TITLE + "\u201D", 12, True, "Signature")
    spacer(doc, 1)
    cover_line(doc, "Submitted by", 12, True, "Signature")
    spacer(doc, 1)
    for s in C.STUDENTS:
        cover_line(doc, f"{s.upper()}     ROLL NO: __________     EXAM NO: __________", 11, True, "Signature")
    spacer(doc, 1)
    p = body_paragraph(
        doc,
        "is a bonafide student of this institute and the work has been carried out by them under the "
        f"supervision of {C.GUIDE} and is approved for the partial fulfillment of the requirement of "
        f"{C.UNIVERSITY}, for the award of the degree of Bachelor of Technology (Information Technology).",
        style="CertificateBody",
    )
    p.paragraph_format.first_line_indent = Inches(0.5)
    spacer(doc, 2)
    cover_line(doc, f"({C.GUIDE})", 12, False, "Signature")
    cover_line(doc, "Guide", 12, False, "Signature")
    spacer(doc, 1)
    cover_line(doc, f"({C.HOD})", 12, False, "Signature")
    cover_line(doc, "Head, Department of Information Technology", 12, False, "Signature")
    spacer(doc, 1)
    cover_line(doc, "Name & Signature of External Examiner: ______________________________", 11, False, "Signature")
    cover_line(doc, f"({C.DIRECTOR})", 12, False, "Signature")
    cover_line(doc, f"Director, {C.COLLEGE}, Pune - 33", 11, False, "Signature")
    spacer(doc, 1)
    cover_line(doc, "Place: Pune                         Date: ____________________", 12, False, "Signature")
    cover_line(doc, "Note: Industry certificate is not applicable; this is an institute-sponsored academic project.", 10, False, "Signature")


def add_acknowledgement(doc: Document) -> None:
    front_heading(doc, "ACKNOWLEDGEMENT", page_break=True)
    body_paragraph(
        doc,
        "We express our sincere gratitude to our project guide, " + C.GUIDE +
        ", for continuous guidance, technical insight and constructive review throughout the project. "
        "The guidance helped us connect hydrological theory, software architecture and research writing into a "
        "single coherent Stage-II submission.",
    )
    body_paragraph(
        doc,
        f"We are grateful to {C.HOD}, Head of the {C.DEPARTMENT}, for providing academic support and "
        f"encouragement, and to {C.DIRECTOR}, Director of {C.COLLEGE}, for the institutional facilities and "
        "environment that made this work possible. We also thank the faculty members of the department for "
        "their valuable suggestions during reviews and evaluations.",
    )
    body_paragraph(
        doc,
        "We acknowledge the open-source communities behind FastAPI, Next.js, ReportLab, HarfBuzz, Leaflet, "
        "react-three-fiber, pytest and the public NOAA PSL CMAP precipitation dataset. Their documented work "
        "enabled a transparent and reproducible implementation. Finally, we thank our families and classmates "
        "for their patience and support during development, testing and report preparation.",
    )
    spacer(doc, 1)
    cover_line(doc, "NAME OF THE STUDENTS", 12, True, "Signature")
    for s in C.STUDENTS:
        cover_line(doc, s.upper() + " (ROLL NO: __________)", 12, True, "Signature")


def add_abstract(doc: Document) -> None:
    front_heading(doc, "ABSTRACT", page_break=True)
    body_paragraph(doc, C.ABSTRACT)
    kp = doc.add_paragraph(style="Normal")
    kp.paragraph_format.first_line_indent = Inches(0)
    kp.paragraph_format.space_before = Pt(6)
    kr = kp.add_run("Keywords: ")
    set_run_font(kr, "Times New Roman", 12, bold=True)
    kr2 = kp.add_run(C.KEYWORDS)
    set_run_font(kr2, "Times New Roman", 12)


def build_toc_entries() -> list[list[str]]:
    entries: list[list[str]] = [
        ["", "CERTIFICATE", FRONT_LABELS["certificate"]],
        ["", "ACKNOWLEDGEMENT", FRONT_LABELS["acknowledgement"]],
        ["", "ABSTRACT", FRONT_LABELS["abstract"]],
        ["", "TABLE OF CONTENTS", FRONT_LABELS["toc1"]],
        ["", "LIST OF ABBREVIATIONS", FRONT_LABELS["abbrev"]],
        ["", "NOMENCLATURE", FRONT_LABELS["nomenclature"]],
        ["", "LIST OF FIGURES", FRONT_LABELS["lof1"]],
        ["", "LIST OF TABLES", FRONT_LABELS["lot1"]],
    ]
    starts: dict[str, int] = {}
    groups: dict[str, list[tuple[int, str]]] = {}
    for i, p in enumerate(PAGES, 1):
        ch = p["ch"]
        starts.setdefault(ch, i)
        groups.setdefault(ch, []).append((i, p["heading"]))
    for ch, items in groups.items():
        if ch.isdigit():
            title = next(p["chapter_title"] for p in PAGES if p["ch"] == ch and p["chapter_title"])
            entries.append([ch, title, str(starts[ch])])
            for i, h in items[:3]:
                entries.append(["", "    " + h, str(i)])
        else:
            title = next(p["chapter_title"] for p in PAGES if p["ch"] == ch and p["chapter_title"])
            entries.append([ch, title, str(starts[ch])])
            entries.append(["", "    " + items[0][1], str(items[0][0])])
    entries.append(["12", "12 REFERENCES", str(REF_START_PAGE)])
    entries.append(["", "    12.3 Bibliography", str(REF_START_PAGE + 2)])
    return entries


def add_toc(doc: Document) -> None:
    entries = build_toc_entries()
    mid = math.ceil(len(entries) / 2)
    front_heading(doc, "TABLE OF CONTENTS", page_break=True)
    add_index_table(doc, ["SR. NO.", "TITLE OF CHAPTER / SECTION", "PAGE NO."], entries[:mid])
    front_heading(doc, "TABLE OF CONTENTS (CONTINUED)", page_break=True)
    add_index_table(doc, ["SR. NO.", "TITLE OF CHAPTER / SECTION", "PAGE NO."], entries[mid:])


def add_abbreviations(doc: Document) -> None:
    front_heading(doc, "LIST OF ABBREVIATIONS", page_break=True)
    items = C.ABBREVIATIONS
    half = math.ceil(len(items) / 2)
    left, right = items[:half], items[half:]
    rows = [["ABBREVIATION", "ILLUSTRATION", "ABBREVIATION", "ILLUSTRATION"]]
    for i in range(half):
        a = left[i]
        b = right[i] if i < len(right) else ("", "")
        rows.append([a[0], a[1], b[0], b[1]])
    add_index_table(doc, rows[0], rows[1:], font_size=8.5)


def add_nomenclature(doc: Document) -> None:
    front_heading(doc, "NOMENCLATURE", page_break=True)
    body_paragraph(
        doc,
        "The following symbols and notations are used in this report. All symbolic definitions are confined to "
        "this section in accordance with the institute report guidelines.",
    )
    rows = [["SYMBOL", "DESCRIPTION", "UNIT"]]
    rows += [list(x) for x in C.NOMENCLATURE]
    add_index_table(doc, rows[0], rows[1:], font_size=8.8)


def build_index_records(pages: list[dict]) -> tuple[list[tuple[str, str, int]], list[tuple[str, str, int]]]:
    fig_records: list[tuple[str, str, int]] = []
    tab_records: list[tuple[str, str, int]] = []
    counters: dict[str, int] = {}
    for page_no, page in enumerate(pages, 1):
        ch = page["ch"]
        counters.setdefault(ch, 0)
        for f in page.get("figures", []):
            counters[ch] += 1
            fig_records.append((f"{ch}.{counters[ch]}", f["caption"], page_no))
        if page.get("table"):
            # Tables use their own sequence, so maintain a separate counter.
            tc = counters.setdefault(ch + "_t", 0)
            tc += 1
            counters[ch + "_t"] = tc
            tab_records.append((f"{ch}.{tc}", page["table"]["caption"], page_no))
    return fig_records, tab_records


def add_list_of_figures(doc: Document, records: list[tuple[str, str, int]]) -> None:
    mid = math.ceil(len(records) / 2)
    front_heading(doc, "LIST OF FIGURES", page_break=True)
    rows = [[n, cap, str(pg)] for n, cap, pg in records[:mid]]
    add_index_table(doc, ["FIGURE", "ILLUSTRATION", "PAGE NO."], rows, font_size=8.5)
    front_heading(doc, "LIST OF FIGURES (CONTINUED)", page_break=True)
    rows = [[n, cap, str(pg)] for n, cap, pg in records[mid:]]
    add_index_table(doc, ["FIGURE", "ILLUSTRATION", "PAGE NO."], rows, font_size=8.5)


def add_list_of_tables(doc: Document, records: list[tuple[str, str, int]]) -> None:
    mid = math.ceil(len(records) / 2)
    front_heading(doc, "LIST OF TABLES", page_break=True)
    rows = [[n, cap, str(pg)] for n, cap, pg in records[:mid]]
    add_index_table(doc, ["TABLE", "ILLUSTRATION", "PAGE NO."], rows, font_size=8.5)
    front_heading(doc, "LIST OF TABLES (CONTINUED)", page_break=True)
    rows = [[n, cap, str(pg)] for n, cap, pg in records[mid:]]
    add_index_table(doc, ["TABLE", "ILLUSTRATION", "PAGE NO."], rows, font_size=8.5)


# ---------------------------------------------------------------------------
# Main report and references
# ---------------------------------------------------------------------------

def add_main_pages(doc: Document, ref_map: dict[str, int],
                   fig_records: list[tuple[str, str, int]],
                   tab_records: list[tuple[str, str, int]]) -> None:
    fig_iter = iter(fig_records)
    tab_iter = iter(tab_records)
    for page_no, page in enumerate(PAGES, 1):
        first = True
        # Chapter / appendix title.
        if page.get("chapter_title"):
            chapter_title(doc, replace_refs(page["chapter_title"], ref_map), page_break=(page_no > 1))
            first = False
        section_heading(doc, replace_refs(page["heading"], ref_map), page_break=(first and page_no > 1))
        first = False
        for para in page["paras"]:
            body_paragraph(doc, replace_refs(para, ref_map), page_break=(first and page_no > 1))
            first = False
        figs = page.get("figures", [])
        for idx, f in enumerate(figs):
            rec = next(fig_iter)
            add_figure(doc, f, rec[0], ref_map, multi=(len(figs) > 1),
                       page_break=(first and page_no > 1))
            first = False
        if page.get("table"):
            rec = next(tab_iter)
            add_table(doc, page["table"]["rows"], page["table"]["caption"], rec[0], ref_map,
                      font_size=8.5 if len(page["table"]["rows"]) > 8 else 9,
                      page_break=(first and page_no > 1))
            first = False
        if page.get("code"):
            add_code_block(doc, page["code"], ref_map, page_break=(first and page_no > 1))
            first = False


def add_references(doc: Document, ref_order: list[str], ref_map: dict[str, int]) -> None:
    refs = [f"[{ref_map[key]}] {C.REFS[key]}" for key in ref_order]
    chunk_size = math.ceil(len(refs) / 3)
    chunks = [refs[:chunk_size], refs[chunk_size:2 * chunk_size], refs[2 * chunk_size:]]
    chapter_title(doc, "12 REFERENCES", page_break=True)
    for text in chunks[0]:
        p = doc.add_paragraph(style="ReferenceEntry")
        p.add_run(text)

    section_heading(doc, "12.1 References (Continued)", page_break=True,
                    style="SubHeading")
    for text in chunks[1]:
        p = doc.add_paragraph(style="ReferenceEntry")
        p.add_run(text)

    section_heading(doc, "12.2 References (Continued)", page_break=True,
                    style="SubHeading")
    for text in chunks[2]:
        p = doc.add_paragraph(style="ReferenceEntry")
        p.add_run(text)

    section_heading(doc, "12.3 Bibliography", style="SubHeading")
    for i, item in enumerate(C.BIBLIOGRAPHY, 1):
        p = doc.add_paragraph(style="ReferenceEntry")
        p.add_run(f"[B{i}] {item}")


def estimate_page_load(pages: list[dict]) -> list[tuple[int, float, str]]:
    """Rough height estimate in inches; used only to flag overflow risk."""
    warnings: list[tuple[int, float, str]] = []
    for i, p in enumerate(pages, 1):
        h = 0.35  # heading allowance
        for para in p["paras"]:
            lines = max(1, math.ceil(len(para) / 88))
            h += lines * 0.25 + 0.18
        nfig = len(p.get("figures", []))
        if nfig:
            h += nfig * (2.08 if nfig > 1 else 3.00)
            h += nfig * 0.42  # caption + legend
        if p.get("table"):
            rows = p["table"]["rows"]
            h += len(rows) * 0.24 + 0.35
        if p.get("code"):
            h += len(p["code"]) * 0.17 + 0.2
        if h > 9.25:
            warnings.append((i, h, p["heading"]))
    return warnings


def build() -> None:
    doc = Document()
    configure_styles(doc)
    set_update_fields(doc)
    doc.core_properties.title = C.TITLE
    doc.core_properties.author = ", ".join(C.STUDENTS)
    doc.core_properties.subject = "BE Project Work - II Stage-II Report"
    doc.core_properties.keywords = C.KEYWORDS

    ref_map, ref_order = build_reference_map(PAGES)
    fig_records, tab_records = build_index_records(PAGES)

    # Section 1 - cover page (footer text, no page number).
    sec1 = doc.sections[0]
    configure_section(sec1)
    configure_footer(sec1, with_page_number=False)
    add_cover(doc)

    # Section 2 - Roman-numbered front matter.
    sec2 = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(sec2, fmt="lowerRoman", start=1)
    configure_footer(sec2, with_page_number=True)
    add_certificate(doc)
    add_acknowledgement(doc)
    add_abstract(doc)
    add_toc(doc)
    add_abbreviations(doc)
    add_nomenclature(doc)
    add_list_of_figures(doc, fig_records)
    add_list_of_tables(doc, tab_records)

    # Section 3 - Arabic-numbered main report from Introduction.
    sec3 = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(sec3, fmt="decimal", start=1)
    configure_footer(sec3, with_page_number=True)
    add_main_pages(doc, ref_map, fig_records, tab_records)
    add_references(doc, ref_order, ref_map)

    doc.save(OUT)
    warnings = estimate_page_load(PAGES)
    print(f"Saved: {OUT}")
    print(f"Main numbered pages (Introduction through appendices): {MAIN_PAGE_COUNT}")
    print(f"Figures: {len(fig_records)}; Tables: {len(tab_records)}; References: {len(ref_order)}")
    print(f"References start at page {REF_START_PAGE}; bibliography page {REF_START_PAGE + 2}")
    if warnings:
        print("Estimated content-height warnings (page, inches, heading):")
        for w in warnings:
            print(" ", w)
    else:
        print("Estimated page loading: no overflow warnings.")


if __name__ == "__main__":
    build()
